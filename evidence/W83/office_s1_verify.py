import tempfile, os
from pathlib import Path
d = Path(tempfile.mkdtemp())
import docx, pptx, openpyxl
from docx import Document
from pptx import Presentation
from openpyxl import Workbook, load_workbook
doc = Document(); doc.add_heading("W83 Test Policy", 1); doc.add_paragraph("Household quiet hours are 10pm to 7am."); doc.save(str(d / "t.docx"))
prs = Presentation()
for title in ("Family Budget 2026", "Savings Goals"):
    s = prs.slides.add_slide(prs.slide_layouts[1]); s.shapes.title.text = title
prs.save(str(d / "t.pptx"))
wb = Workbook(); wb.active["A1"] = "Groceries"; wb.active["B1"] = 450; wb.save(str(d / "t.xlsx"))
r1 = Document(str(d / "t.docx")); r2 = Presentation(str(d / "t.pptx")); r3 = load_workbook(str(d / "t.xlsx")).active
print("RT docx heading=%r paras=%d" % (r1.paragraphs[0].text, len(r1.paragraphs)))
print("RT pptx slides=%d titles=%s" % (len(r2.slides), [s.shapes.title.text for s in r2.slides]))
print("RT xlsx A1=%r B1=%r" % (r3["A1"].value, r3["B1"].value))
from openjarvis.server.upload_router import _extract_text_from_docx
txt = _extract_text_from_docx((d / "t.docx").read_bytes())
print("G-4 author docx reader returns text=%s content=%r" % ("quiet hours" in txt, txt[:80]))
print("VERSIONS python-docx=%s python-pptx=%s openpyxl=%s" % (getattr(docx, "__version__", "?"), pptx.__version__, openpyxl.__version__))
import shutil; shutil.rmtree(d, ignore_errors=True)
