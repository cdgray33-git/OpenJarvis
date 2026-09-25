# s3_family_office_w87.py - W87 S3 re-run 4 (live V&V of patch A openjarvis-w87-codefiles-v1)
# Same 6 family requests and detail checks as evidence\W83\s3_family_office.py, plus per request:
#   "reply honest" - if the file exists and opens, the reply must not claim failure (R2 failure class)
#   shell_exec attempt count (the gated self-check path)
# Unopenable or missing file fails EVERY check of that request. Writes evidence\W83\s3-rerun4.txt and prints it.
# Run from PS C:\Users\Admin\OpenJarvis> with the venv python.
import json, time, glob, os, re, sys
from pathlib import Path
import httpx

OUT = Path.cwd() / "evidence" / "W83" / "s3-rerun4.txt"
_lines = []
def say(s=""):
    print(s); _lines.append(s)

B = "http://127.0.0.1:8010"; WS = Path.home() / ".openjarvis" / "workspace"; DL = Path(os.environ["LOCALAPPDATA"]) / "OpenJarvis" / "logs" / "dispatch.log"
REQ = [
 ("r1", "w83_s3_r1_vacation.pptx", "Make a PowerPoint presentation for my family about our summer vacation plan. It must have exactly 5 slides: a title slide titled 'Summer Vacation 2026', then one slide each titled Budget, Schedule, Packing, and Chores, each with a few bullet points. Save it as w83_s3_r1_vacation.pptx."),
 ("r2", "w83_s3_r2_screentime.docx", "Write a household screen-time policy as a Word document with three sections headed Purpose, Rules by Age, and Consequences. Make it at least 200 words. Save it as w83_s3_r2_screentime.docx."),
 ("r3", "w83_s3_r3_budget.xlsx", "Create an Excel monthly household budget with the columns Category, Budgeted, Actual, and Difference. Include rows for Rent, Groceries, Utilities, Transportation, and Savings, and a Total row. The Difference column must use formulas. Save it as w83_s3_r3_budget.xlsx."),
 ("r4", "w83_s3_r4_safety.pptx", "Make a 3-slide PowerPoint for my kids about internet safety rules: a title slide and two content slides, each with at least 3 bullet points. Save it as w83_s3_r4_safety.pptx."),
 ("r5", "w83_s3_r5_vehicle.docx", "Write a Family Vehicle Use Policy as a Word document with sections headed Scope, Eligibility, Rules, Fuel and Maintenance, and Violations. The Rules section must be a numbered list of at least 5 rules. Save it as w83_s3_r5_vehicle.docx."),
 ("r6", "w83_s3_r6_letter.docx", "Write a Word document letter to my daughter's school principal requesting a meeting about her progress. Start with 'Dear Principal', state the purpose, propose two specific meeting times, and close with 'Sincerely'. Save it as w83_s3_r6_letter.docx.")]
FAILWORDS = ["wasn't able", "was not able", "not able to", "unable", "could not", "couldn't", "failed", "cannot create", "can't create", "was not created", "wasn't created"]

def check(rid, f):
    c = []
    if rid in ("r1", "r4"):
        from pptx import Presentation
        p = Presentation(str(f)); sl = list(p.slides)
        tt = [(s.shapes.title.text if s.shapes.title is not None else "") for s in sl]
        def bullets(s): return sum(1 for sh in s.shapes if sh.has_text_frame and sh != s.shapes.title for para in sh.text_frame.paragraphs if para.text.strip())
        if rid == "r1":
            c += [("5 slides", len(sl) == 5), ("title Summer Vacation", "summer vacation" in (tt[0].lower() if tt else ""))]
            c += [("slide " + w, any(w.lower() in t.lower() for t in tt)) for w in ("Budget", "Schedule", "Packing", "Chores")]
        else:
            c += [("3 slides", len(sl) == 3)] + [("slide %d >=3 bullets" % (i + 2), len(sl) > i + 1 and bullets(sl[i + 1]) >= 3) for i in range(2)]
    elif rid == "r3":
        from openpyxl import load_workbook
        ws = load_workbook(str(f)).active; cells = [str(v) for row in ws.iter_rows(values_only=True) for v in row if v is not None]; low = [x.lower() for x in cells]
        c += [("header " + h, any(h.lower() == x.strip() for x in low)) for h in ("Category", "Budgeted", "Actual", "Difference")]
        c += [("row " + r, any(r.lower() in x for x in low)) for r in ("Rent", "Groceries", "Utilities", "Transportation", "Savings")]
        c += [("Total row", any("total" in x for x in low)), ("formulas", any(x.startswith("=") for x in cells))]
    else:
        from docx import Document
        d = Document(str(f)); paras = [p.text for p in d.paragraphs]; text = "\n".join(paras); low = text.lower()
        if rid == "r2":
            c += [("section " + h, h.lower() in low) for h in ("Purpose", "Rules by Age", "Consequences")] + [(">=200 words", len(text.split()) >= 200)]
        elif rid == "r5":
            c += [("section " + h, h.lower() in low) for h in ("Scope", "Eligibility", "Rules", "Fuel and Maintenance", "Violations")]
            items = sum(1 for p in d.paragraphs if p.text.strip() and ("List" in (p.style.name or "") or re.match(r"^\s*\d+[.)]", p.text)))
            c += [(">=5 numbered rules", items >= 5)]
        else:
            times = re.findall(r"\b\d{1,2}(:\d{2})?\s*(am|pm|a\.m\.|p\.m\.)", low)
            c += [("Dear Principal", "dear principal" in low), ("meeting", "meeting" in low), (">=2 times", len(times) >= 2), ("Sincerely", "sincerely" in low)]
    return c

EXPECTED = {"r1": 7, "r2": 5, "r3": 12, "r4": 4, "r5": 7, "r6": 5}
say("W87 Office S3 RE-RUN 4 (patch A openjarvis-w87-codefiles-v1, before patch B) %s" % time.strftime("%Y-%m-%dT%H:%M:%S"))
WS.mkdir(parents=True, exist_ok=True)
for g in WS.glob("w83_s3_*"): g.unlink()
root_before = set(glob.glob("*.docx") + glob.glob("*.pptx") + glob.glob("*.xlsx"))
tot_p = tot_n = honest = shells = 0
for rid, fname, q in REQ:
    n0 = len(DL.read_text(errors="replace").splitlines()); t0 = time.time()
    try:
        r = httpx.post(B + "/v1/chat/completions", json={"model": "qwen3-coder:30b", "stream": False, "max_tokens": 2048, "messages": [{"role": "user", "content": q}]}, timeout=300)
        reply = r.json()["choices"][0]["message"]["content"] or ""
    except Exception as e:
        reply = "CHAT ERR %s" % e
    dur = time.time() - t0
    oc = [l for l in DL.read_text(errors="replace").splitlines()[n0:] if "OUTCOME" in l]
    ci = [l for l in oc if "tool=code_interpreter" in l]
    sx = sum(1 for l in oc if "tool=shell_exec" in l); shells += sx
    f = WS / fname
    if f.exists():
        try:
            c = check(rid, f)
            rl = reply.lower(); bad = [w for w in FAILWORDS if w in rl]
            c.append(("reply honest" + (" (claims failure: %s)" % ", ".join(bad) if bad else ""), not bad))
        except Exception as e:
            c = [("file opens", False)] * EXPECTED[rid]; say("   CHECK ERROR %s: %s" % (rid, e))
    else:
        c = [("file created", False)] * EXPECTED[rid]
    if any(n.startswith("reply honest") and ok for n, ok in c): honest += 1
    p = sum(1 for _, ok in c if ok); tot_p += p; tot_n += len(c)
    say("%s %-26s file=%-5s score=%d/%d  dur=%ds  code_interpreter_calls=%d ok=%d  shell_exec=%d" % (rid.upper(), fname, f.exists(), p, len(c), dur, len(ci), sum(1 for l in ci if "success=True" in l), sx))
    for name, ok in c: say("     %s %s" % ("PASS" if ok else "FAIL", name))
    for l in oc: say("     " + l[l.find("tool="):][:120])
    say("     reply: " + re.sub(r"[^\x20-\x7E]", " ", reply)[:300])
new_root = set(glob.glob("*.docx") + glob.glob("*.pptx") + glob.glob("*.xlsx")) - root_before
real = sum(1 for _, fn, _ in REQ if (WS / fn).exists())
say("TOTAL detail score %d/%d   files present %d/6   honest replies %d/6   shell_exec attempts %d" % (tot_p, tot_n, real, honest, shells))
say("RQ-032 CRITERION (6/6 real files AND 6/6 honest replies): %s" % ("MET" if real == 6 and honest == 6 else "NOT MET"))
say("G-1 new Office files in repo root: %s" % (sorted(new_root) or "none"))
OUT.write_text("\n".join(_lines) + "\n", encoding="utf-8")
print("written:", OUT)
